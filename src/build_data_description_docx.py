"""Build an editable, student-voiced Data in Brief-style manuscript."""

from pathlib import Path
import csv

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "output" / "docx" / "chirundu_cdf_data_description_paper_editable.docx"
FILES = [p for p in sorted((ROOT / "data" / "processed").rglob("*.csv")) if "cdf_pilot" not in p.parts]
assert len(FILES) == 10

def row_count(path):
    with path.open(encoding="utf-8-sig", newline="") as stream:
        return sum(1 for _ in csv.DictReader(stream, delimiter="|"))

assert sum(map(row_count, FILES)) == 730

DOC = Document()
section = DOC.sections[0]
section.page_width = Inches(8.27)
section.page_height = Inches(11.69)
section.top_margin = Inches(0.83)
section.bottom_margin = Inches(0.73)
section.left_margin = Inches(0.83)
section.right_margin = Inches(0.83)
section.header_distance = Inches(0.35)
section.footer_distance = Inches(0.38)

BLACK = RGBColor(0, 0, 0)
MUTED = RGBColor(70, 70, 70)
LIGHT_BLUE = "EAF1F5"
BORDER = "D9D9D9"

def style(name, size, bold=False, before=0, after=0, line=1.16, keep_next=False):
    s = DOC.styles[name]
    s.font.name = "Aptos"
    s.font.size = Pt(size)
    s.font.bold = bold
    s.font.color.rgb = BLACK
    s.paragraph_format.space_before = Pt(before)
    s.paragraph_format.space_after = Pt(after)
    s.paragraph_format.line_spacing = line
    s.paragraph_format.keep_with_next = keep_next
    return s

style("Normal", 10.5, after=6, line=1.16)
style("Title", 16, bold=True, after=10, line=1.08, keep_next=True)
DOC.styles.add_style("Byline", WD_STYLE_TYPE.PARAGRAPH)
style("Byline", 9.5, after=12, line=1.16, keep_next=True)
style("Heading 1", 12, bold=True, before=12, after=5, keep_next=True)
style("Heading 2", 10.5, bold=True, before=8, after=4, keep_next=True)
style("Caption", 9, bold=True, before=8, after=4, keep_next=True)
for styled in (DOC.styles["Title"], DOC.styles["Byline"]):
    ppr = styled.element.get_or_add_pPr()
    for border in ppr.findall(qn("w:pBdr")):
        ppr.remove(border)

def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)

def borders(cell):
    tc_pr = cell._tc.get_or_add_tcPr()
    b = OxmlElement("w:tcBorders")
    for side in ("top", "left", "bottom", "right"):
        e = OxmlElement(f"w:{side}")
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), "4")
        e.set(qn("w:color"), BORDER)
        b.append(e)
    tc_pr.append(b)

def cell_margin(cell, top=75, start=90, bottom=75, end=90):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    mar = OxmlElement("w:tcMar")
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:w"), str(value))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tc_pr.append(mar)

def table(rows, widths, header=False):
    t = DOC.add_table(rows=len(rows), cols=len(rows[0]))
    t.autofit = False
    for i, row in enumerate(rows):
        for j, value in enumerate(row):
            c = t.cell(i, j)
            c.width = Inches(widths[j])
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            c.text = str(value)
            borders(c)
            cell_margin(c)
            if (header and i == 0) or (not header and j == 0):
                shade(c, LIGHT_BLUE)
            elif header and i % 2 == 0:
                shade(c, "F8FAFB")
            for p in c.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.08
                for run in p.runs:
                    run.font.name = "Aptos"
                    run.font.size = Pt(8.5 if header else 8.7)
                    run.font.color.rgb = BLACK
                    if (header and i == 0) or (not header and j == 0):
                        run.font.bold = True
                if header and j == 1:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if header:
        for c in t.rows[0].cells:
            c._tc.get_or_add_tcPr().append(OxmlElement("w:tblHeader"))
    DOC.add_paragraph().paragraph_format.space_after = Pt(0)
    return t

def add(text, style_name=None):
    return DOC.add_paragraph(text, style_name)

def h1(text):
    p = add(text, "Heading 1")
    if text in ("Data Description", "Data Access and Reuse"):
        p.paragraph_format.page_break_before = True

def h2(text):
    add(text, "Heading 2")

def bullet(text):
    p = add(text, "Normal")
    p.style = DOC.styles["List Bullet"]
    p.paragraph_format.left_indent = Inches(0.22)
    p.paragraph_format.first_line_indent = Inches(-0.14)
    p.paragraph_format.space_after = Pt(5)

def page_number(paragraph):
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    paragraph._p.append(field)

head = section.header.paragraphs[0]
head.text = "CSC 4792  |  DATA DESCRIPTION PAPER"
head.style = DOC.styles["Normal"]
head.paragraph_format.space_after = Pt(0)
for r in head.runs:
    r.font.size = Pt(8)
    r.font.color.rgb = MUTED
foot = section.footer.paragraphs[0]
foot.alignment = WD_ALIGN_PARAGRAPH.RIGHT
page_number(foot)

add("Chirundu Town Council CDF and local finance data 2022 to 2026", "Title")
add("CSC 4792 Data Mining and Warehousing  |  Group 39  |  University of Zambia\nPrepared 14 September 2026", "Byline")

h1("Abstract")
add("We created ten CSV tables from 16 PDF documents published by Chirundu Town Council. Together, the tables contain 730 rows on approved Constituency Development Fund (CDF) community projects, council budgets, revenue estimates, financial statements and selected performance reports. The records cover project lists from 2022 to 2026, signed financial statements from 2022 to 2024, a 2023 procurement report and 2025 half-year reporting. We extracted selectable text directly and used optical character recognition for scanned pages. We then checked difficult entries against the original PDF images and kept a source URL and page number with every row. This paper explains how we made the tables, what each one contains and where readers need to be careful when reusing them. The CSV files and notebook are in our course project. We could not verify a public repository link when preparing this version of the paper.")
add("Keywords  Chirundu; Zambia; Constituency Development Fund; local government; public finance; administrative data")

h1("Specifications table")
table([
    ("Subject", "Public administration and local government finance"),
    ("Specific subject area", "CDF projects, council budgets and performance reporting in Chirundu, Zambia"),
    ("Type of data", "Ten structured CSV tables from council PDF documents"),
    ("Data collection", "We downloaded council PDFs, extracted selectable text or used OCR for scanned tables, then checked uncertain cells against page images."),
    ("Data format", "UTF-8 CSV files with pipe (|) separators; blank cells mean unavailable information"),
    ("Data source location", "Chirundu Town Council, Chirundu District, Zambia; https://www.chirunducouncil.gov.zm"),
    ("Data accessibility", "The ten CSV files are in the accompanying course project under data/processed/. We did not have a verified public dataset URL or persistent identifier at preparation."),
    ("Related work", "Our extraction notebook is notebooks/chirundu_cdf_project.ipynb"),
], [1.52, 5.08])

h1("Value of the Data")
bullet("The tables bring several years of council project and finance records into a consistent format while keeping links to the pages they came from.")
bullet("Other students, researchers and civic groups can use the files to examine what information the council published and compare it with similar records from other councils.")
bullet("We kept approved projects, budgets, financial statements and progress reports separate, so readers can compare them without confusing planned amounts with actual spending.")
bullet("The source links, notes and notebook make it easier for others to check our transcription, correct it and add later reports.")

h1("Data Description")
add("Table 1 lists the ten files we describe in this paper. Every name begins with db-unza26-csc4792-chirundu_, as required for the assignment. All files have a header row and use the pipe character as the separator. We count rows as entries in the source tables; 730 rows do not mean 730 separate projects or payments. We kept two earlier 2025 pilot files for comparison with our extraction process, but left them out of this release because their project entries appear in the main approved-project file.")
add("Table 1  Curated data files and coverage", "Caption")
table([
    ("File suffix  csv", "Rows", "What one row describes"),
    ("cdf_approved_projects_2022_2026", "83", "An entry in an annual approved-project list, 2022 to 2026"),
    ("budget_programmes", "70", "A printed programme allocation, 2022 to 2026"),
    ("cdf_budget_subprogrammes", "22", "A CDF subprogramme allocation, 2022 to 2026"),
    ("budget_revenue", "231", "A detailed revenue estimate, 2022, 2023, 2025 or 2026"),
    ("budget_totals", "5", "A printed annual council budget total, 2022 to 2026"),
    ("financial_statements", "207", "A line in a signed financial statement, 2022 to 2024"),
    ("budget_vs_actual", "66", "A final budget and actual comparison line, 2022 to 2024"),
    ("cdf_project_progress_2023_q1", "8", "A procurement or contract observation in early 2023"),
    ("cdf_performance_indicators_2025", "7", "A CDF indicator reported to 30 June 2025"),
    ("half_year_finances_2025", "31", "A finance line for January to June 2025"),
    ("Total", "730", "Ten distinct files"),
], [2.45, 0.48, 3.67], header=True)

add("The approved-project file records the project number, list year, description, reported sector and location, approval wording, any stated amount, and a source URL and page. The year-by-year counts are 16, 19, 15, 14 and 19 for 2022 through 2026. A project can appear in more than one year's list, so we have not treated the 83 entries as 83 unique physical projects. Only one approved-list entry gives an explicit monetary amount. The lists do not report implementation progress, so we left all 83 implementation_status cells blank rather than guessing.")
add("The budget tables contain programme, subprogramme and revenue lines in Zambian kwacha (ZMW). Financial records identify the statement, section and row type, which helps readers distinguish details from subtotals and totals. The performance tables describe reporting periods, indicators or procurement stages. Six of the eight 2023 progress entries could be linked to 2022 approvals after we compared locations and work descriptions. Those links are our documented matches, not official project identifiers.")

h1("Experimental Design Materials and Methods")
h2("Source selection and collection")
add("We collected documents from Chirundu Town Council's public website. Our download inventory contains 104 files covering CDF, finance, performance, procurement, planning, governance and other council topics. For this release, we selected 16 PDFs: five annual approved-project lists, five annual budgets, three signed financial reports, two 2025 half-year reports and one 2023 procurement report. We did not add duplicate copies or alternative project lists as extra records. The five selected approved-project PDFs are inventory items 001, 002, 003, 059 and 005. In each output row, source_url identifies the council PDF and source_page gives the page number starting from the first PDF page. We took the year from the document heading rather than the upload date.")
h2("Extraction and cleaning")
add("Our Jupyter notebook shows the extraction steps. For PDFs with selectable text, we used pdfplumber. For scanned pages, we rendered the pages as images and used RapidOCR to read them. We used the positions of words and table columns to group the OCR output into rows. One 2023 project-list page needed rotation before OCR. The notebook walks through the 2025 approved-project list in detail, and we kept raw OCR output and intermediate page data in data/interim/.")
add("We checked unclear words, amounts and column boundaries against the original images, especially where stamps or borders interfered with recognition. We joined descriptions that wrapped across lines, cleaned spacing, removed duplicate extracted rows and parsed monetary values without their thousands separators. We did not turn blanks or dashes into zeroes. We also kept notes where source text was obscured, a label needed restoration or printed figures disagreed. The README files beside the CSVs explain the fields and table-specific exceptions.")
h2("Variables and checks")
add("The files can be compared by year and source, but they are not one transaction ledger. A project_no only identifies a row within one year's approved list. Reporting start and end dates in the performance files are different from project approval years. In budget_vs_actual, we calculated variance as actual_amount_zmw minus final_budget_zmw. If either value is missing, the variance is blank. The 2025 half-year report prints its variance in the opposite direction, budget minus actual, so we kept it in a separate reported_variance_zmw field. We also retained percentage fields as text when the source displayed formula errors or implausible values.")
add("We checked that the ten files could be parsed with the pipe separator, that their row counts added to 730, and that all 730 rows contained a source URL. We did not force small differences in published figures to agree. For example, the 2025 CDF programme allocation differs by K1 from the sum of its subprogrammes, and some financial statement lines differ by K1 between related tables. A 2023 statement also prints the wrong year in a closing-balance label. The 2025 half-year report contains formula artefacts such as #DIV/0!. We kept these issues visible in the data and notes.")
h2("Limits on interpretation")
add("This is a selected release from a larger collection of council documents. We have not extracted every council function, every CDF category or project-level disbursements. The selected 2024 budget has no detailed revenue schedule, and our release has no 2025 or 2026 annual actual financial statements. An approved project is not necessarily completed. A procurement estimate or contract amount is not an actual payment. Readers also need to avoid adding detail rows to their subtotals, or council totals to fund totals, because that would count the same money twice. Ward names vary in spelling between source documents, so ward-level comparisons need another review.")

h1("Data Access and Reuse")
add("We placed the ten CSV files in data/processed/cdf_projects/, data/processed/finance/ and data/processed/performance/ in the course project. Our notebook and the intermediate OCR files show how we made them. Each CSV also points back to the source PDF and page. We could not verify a public Kaggle link or other persistent identifier in the project materials when writing this document. Anyone using this version should work from the accompanying files; we will need to add the public repository link when the dataset has been deposited.")
add("The source documents were published by Chirundu Town Council, and we ask anyone reusing our tables to cite the specific council PDFs shown in the rows they use. This release describes public project and financial reports. We did not include the unextracted beneficiary lists in these ten tables. Any later work with individual-level records would need a separate review before publication.")

h1("Declaration of Interests")
add("The project materials do not contain a group declaration of interests. Each group member should confirm this section before an external journal submission.")

h1("References")
add("[1] Chirundu Town Council. Official website and published PDF documents. https://www.chirunducouncil.gov.zm. Source-specific URLs and page numbers are in the CSV files.")
add("[2] CSC 4792 Data Mining and Warehousing. Mini Project Practical Assignment. University of Zambia, 4 September 2026. Assignment brief supplied with the project.")
add("[3] Wang, H.-R. How to write a good Data in Brief article. Elsevier Researcher Academy, June 2024. https://researcheracademy.elsevier.com/uploads/2024-07/Quick_guide_Data_in_brief.pdf.")
add("[4] Data in Brief. Guide for Authors. https://www.sciencedirect.com/journal/data-in-brief/publish/guide-for-authors. Accessed 14 September 2026.")

OUT.parent.mkdir(parents=True, exist_ok=True)
DOC.save(OUT)
print(OUT)
